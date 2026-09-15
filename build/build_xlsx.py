# -*- coding: utf-8 -*-
import sys
from pathlib import Path
BUILD_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(BUILD_DIR))
from investor_data import SECTIONS, TSS
from outreach_messages import MESSAGES
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

NAVY = "0D2A4D"
STEEL = "4682B4"
GOLD = "C9A84C"
LIGHT = "EEF3F8"
SECTIONFILL = "2D5378"

wb = Workbook()

# ---------- Sheet 1: Investor Database ----------
ws = wb.active
ws.title = "Investor Database"

headers = ["Investor Type", "Firm / Name", "Contact Name", "Title", "Geography",
           "Primary Sectors", "Stage Focus", "Typical Check Size", "Website",
           "Email", "Phone", "LinkedIn", "Notes",
           "Outreach \u2013 Formal", "Outreach \u2013 Conversational", "Outreach \u2013 Bold"]

hdr_fill = PatternFill("solid", fgColor=NAVY)
hdr_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
cell_font = Font(name="Arial", size=9)
sec_fill = PatternFill("solid", fgColor=SECTIONFILL)
sec_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
thin = Side(style="thin", color="C9D4E0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
top_align = Alignment(vertical="top", wrap_text=True)

# header row
ws.append(headers)
for c in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=c)
    cell.fill = hdr_fill
    cell.font = hdr_font
    cell.alignment = Alignment(vertical="center", wrap_text=True, horizontal="center")
    cell.border = border
ws.row_dimensions[1].height = 30

r = 2
for sec in SECTIONS:
    # section separator row
    ws.cell(row=r, column=1, value=sec["name"])
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(headers))
    sc = ws.cell(row=r, column=1)
    sc.fill = sec_fill
    sc.font = sec_font
    sc.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[r].height = 22
    r += 1
    msgs = MESSAGES.get(sec["id"], {})
    for inv in sec["investors"]:
        firm, contact, title, geo, sectors, stage, check, website, email, phone, notes = inv
        row = [sec["name"], firm, contact, title, geo, sectors, stage, check,
               website, email, phone, "", notes,
               msgs.get("formal", ""), msgs.get("conversational", ""), msgs.get("bold", "")]
        ws.append(row)
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = cell_font
            cell.alignment = top_align
            cell.border = border
        r += 1

# column widths
widths = [30, 30, 20, 22, 22, 30, 16, 18, 26, 22, 16, 20, 40, 55, 55, 55]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.freeze_panes = "A2"
ws.auto_filter.ref = "A1:" + get_column_letter(len(headers)) + "1"

# ---------- Sheet 2: Outreach Messages ----------
ws2 = wb.create_sheet("Outreach Messages")
ws2.append(["Investor Type", "Tone", "Message (personalise [First Name] / [Firm] before sending)"])
for c in range(1, 4):
    cell = ws2.cell(row=1, column=c)
    cell.fill = hdr_fill
    cell.font = hdr_font
    cell.alignment = Alignment(vertical="center", wrap_text=True, horizontal="center")
    cell.border = border
ws2.row_dimensions[1].height = 26
rr = 2
tone_labels = [("formal", "Formal / Corporate"), ("conversational", "Conversational Professional"), ("bold", "Bold / Disruptive")]
for sec in SECTIONS:
    msgs = MESSAGES.get(sec["id"], {})
    for key, label in tone_labels:
        ws2.cell(row=rr, column=1, value=sec["name"])
        ws2.cell(row=rr, column=2, value=label)
        ws2.cell(row=rr, column=3, value=msgs.get(key, ""))
        for c in range(1, 4):
            cell = ws2.cell(row=rr, column=c)
            cell.font = cell_font
            cell.alignment = top_align
            cell.border = border
        rr += 1
ws2.column_dimensions["A"].width = 34
ws2.column_dimensions["B"].width = 26
ws2.column_dimensions["C"].width = 100
ws2.freeze_panes = "A2"

# ---------- Sheet 3: Methodology ----------
ws3 = wb.create_sheet("Methodology & Notes")
title_font = Font(name="Arial", bold=True, color=NAVY, size=14)
sub_font = Font(name="Arial", bold=True, color=NAVY, size=11)
body_font = Font(name="Arial", size=10)
ws3.column_dimensions["A"].width = 110
def add(txt, font, height=None):
    global rr3
    ws3.cell(row=rr3, column=1, value=txt).font = font
    ws3.cell(row=rr3, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    if height:
        ws3.row_dimensions[rr3].height = height
    rr3 += 1
rr3 = 1
add("The Secret Server \u2014 Investor Database (Wave 1)", title_font, 22)
add("", body_font)
add("About the raise", sub_font)
add("The Secret Server is " + TSS["one_liner"] + ". We are raising our first institutional "
    "round (seed, positionable as Series A) to accelerate a live UK and European partnership across "
    "US, Canadian, UK and EU markets.", body_font, 60)
add("", body_font)
add("How this list was built", sub_font)
add("Every entry is a real, publicly known firm, fund or angel network selected for relevance to "
    "hospitality technology, AR/VR, enterprise SaaS, edtech/workforce training, and consumer/growth "
    "investing across the four target geographies. Firms are grouped by investor type, and each type has "
    "three ready-to-personalise outreach messages (Formal, Conversational, Bold).", body_font, 60)
add("", body_font)
add("A note on contact details", sub_font)
add("Direct personal emails and phone numbers for investment partners are frequently not published and "
    "cannot be verified from public sources. Where they are not reliably known, those fields are left blank "
    "rather than guessed \u2014 no contact information has been fabricated. Each firm's website is the reliable "
    "contact anchor: team pages, pitch-submission forms and general enquiry addresses live there.", body_font, 75)
add("", body_font)
add("Recommended approach", sub_font)
add("Warm introductions convert far better than cold outreach. Where possible, seek a mutual connection "
    "(LinkedIn, portfolio founders, advisers) before sending a cold message. When going direct, use each "
    "firm's submission form or general contact, personalise the [First Name] and [Firm] merge fields, and "
    "lead with the strategic fit noted for that investor type.", body_font, 75)
add("", body_font)
add("This is Wave 1 of an expanding database \u2014 further waves will add deeper regional and sector coverage.", sub_font, 30)
add("Contact: " + TSS["contact_name"] + ", " + TSS["contact_role"] + "  \u00b7  " + TSS["phone"] + "  \u00b7  " + TSS["email"] + "  \u00b7  " + TSS["web"], body_font, 20)

out = BUILD_DIR.parent / "TSS_Investor_Database_Wave1.xlsx"
wb.save(out)

# counts
total = sum(len(s["investors"]) for s in SECTIONS)
print("Saved:", out)
print("Sections:", len(SECTIONS), "Total investors:", total)
for s in SECTIONS:
    print("  -", s["id"], len(s["investors"]))
